"""Workbook parsing.

``openpyxl`` is used when it is installed. Otherwise a small reader built on
the standard library handles the common institutional workbook: shared and
inline strings, numbers, booleans, and ISO dates. Either way a %-formatted cell
is read as the percentage it shows (85, not Excel's stored 0.85). Legacy
``.xls`` binaries need the optional ``xlrd`` package.
"""

from __future__ import annotations

import re
import zipfile
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any
from xml.etree import ElementTree

from ..models import FileKind, ParseResult, ParsedTable, ParserError, ParserUnavailable
from .archive import check_archive_limits, open_entry, read_entry
from .tabular import MAX_ROWS, grid_to_table

_NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
_CELL_REF = re.compile(r"([A-Z]+)(\d+)")
_EXCEL_EPOCH = datetime(1899, 12, 30)
# Built-in number formats Excel treats as dates (ECMA-376 §18.8.30).
_DATE_FORMAT_IDS = {14, 15, 16, 17, 18, 19, 20, 21, 22, 27, 28, 29, 30, 31, 32, 33, 34, 35, 36, 45, 46, 47, 50, 51, 52, 53, 54, 55, 56, 57, 58}
# Built-in percent formats: 0% and 0.00%.
_PERCENT_FORMAT_IDS = {9, 10}
# Parts of a number format where a % sign does not scale the number: escaped
# (\%) or padding (_% *%) characters, quoted text, and [..] sections.
_LITERAL_FORMAT_PARTS = re.compile(r'\\.|[_*].|"[^"]*"|\[[^\]]*\]')


# Cells beyond this column are ignored: no institutional data table needs them,
# and padding rows out to a crafted reference such as ZZZZZZ1 would allocate
# hundreds of millions of cells.
MAX_SHEET_COLUMNS = 1024


def _column_index(ref: str) -> int:
    match = _CELL_REF.match(ref)
    if not match:
        return 0
    letters = match.group(1)
    if len(letters) > 3:  # Excel's last column is XFD
        return MAX_SHEET_COLUMNS
    index = 0
    for char in letters:
        index = index * 26 + (ord(char) - ord("A") + 1)
    return index - 1


def _excel_serial_to_date(value: float) -> Any:
    try:
        result = _EXCEL_EPOCH + timedelta(days=float(value))
    except (OverflowError, ValueError):
        return value
    if result.hour == 0 and result.minute == 0 and result.second == 0:
        return result.date().isoformat()
    return result.isoformat()


def _is_percent_format(code: str) -> bool:
    return "%" in code and "%" in _LITERAL_FORMAT_PARTS.sub("", code)


def _percent_value(number: float) -> Any:
    """The percentage a %-formatted cell shows; Excel stores 85% as 0.85."""

    # Moving the decimal point on the digits avoids float noise (0.07 * 100 is 7.000000000000001).
    shown = float(Decimal(repr(number)).scaleb(2))
    if 0 < shown < 1:
        # The percent cleaner takes a bare number between 0 and 1 for a fraction
        # (0.5 -> 50%), so a cell under 1% keeps its percent sign.
        return f"{shown}%"
    return int(shown) if shown.is_integer() else shown


def _shared_strings(archive: zipfile.ZipFile) -> list[str]:
    """Stream the shared string table; it can be as large as the sheets themselves."""

    try:
        source = open_entry(archive, "xl/sharedStrings.xml")
    except KeyError:
        return []
    strings: list[str] = []
    item_tag = f"{{{_NS['m']}}}si"
    text_tag = f"{{{_NS['m']}}}t"
    try:
        with source:
            for _, element in ElementTree.iterparse(source, events=("end",)):
                if element.tag != item_tag:
                    continue
                strings.append("".join(node.text or "" for node in element.iter(text_tag)))
                element.clear()
    except (ElementTree.ParseError, zipfile.BadZipFile) as exc:
        raise ParserError("workbook shared strings could not be read") from exc
    return strings


def _read_xml(archive: zipfile.ZipFile, name: str) -> ElementTree.Element:
    """Parse one small part of the package; ``KeyError`` when it is absent."""

    try:
        return ElementTree.fromstring(read_entry(archive, name))
    except ElementTree.ParseError as exc:
        raise ParserError(f"workbook part {name} is not well-formed XML") from exc


def _number_styles(archive: zipfile.ZipFile) -> tuple[set[int], set[int]]:
    """The cell styles that show a date, and those that show a percentage."""

    try:
        root = _read_xml(archive, "xl/styles.xml")
    except KeyError:
        return set(), set()
    custom_date_formats: set[int] = set()
    custom_percent_formats: set[int] = set()
    for fmt in root.iterfind("m:numFmts/m:numFmt", _NS):
        code = (fmt.get("formatCode") or "").lower()
        stripped = re.sub(r"\[[^\]]*\]|\"[^\"]*\"", "", code)
        try:
            fmt_id = int(fmt.get("numFmtId", "-1"))
        except ValueError:
            continue
        if re.search(r"[dmy]", stripped) and not re.search(r"[#0]", stripped):
            custom_date_formats.add(fmt_id)
        elif _is_percent_format(code):
            custom_percent_formats.add(fmt_id)
    date_styles: set[int] = set()
    percent_styles: set[int] = set()
    for index, xf in enumerate(root.iterfind("m:cellXfs/m:xf", _NS)):
        try:
            fmt_id = int(xf.get("numFmtId", "0"))
        except ValueError:
            continue
        if fmt_id in _DATE_FORMAT_IDS or fmt_id in custom_date_formats:
            date_styles.add(index)
        elif fmt_id in _PERCENT_FORMAT_IDS or fmt_id in custom_percent_formats:
            percent_styles.add(index)
    return date_styles, percent_styles


def _sheet_paths(archive: zipfile.ZipFile) -> list[tuple[str, str]]:
    try:
        workbook = _read_xml(archive, "xl/workbook.xml")
        rels = _read_xml(archive, "xl/_rels/workbook.xml.rels")
    except KeyError as exc:
        raise ParserError("workbook is missing its sheet index (xl/workbook.xml)") from exc
    targets = {rel.get("Id"): rel.get("Target", "") for rel in rels.findall("rel:Relationship", _NS)}
    sheets: list[tuple[str, str]] = []
    for sheet in workbook.iterfind("m:sheets/m:sheet", _NS):
        rel_id = sheet.get(f"{{{_NS['r']}}}id")
        target = targets.get(rel_id, "")
        if not target:
            continue
        # A target starting with / is relative to the package root, not to xl/.
        path = target.lstrip("/") if target.startswith("/") else target if target.startswith("xl/") else f"xl/{target}"
        if sheet.get("state") == "hidden":
            continue
        sheets.append((sheet.get("name", f"Sheet{len(sheets) + 1}"), path))
    return sheets


def _cell_value(cell: ElementTree.Element, shared: list[str], date_styles: set[int], percent_styles: set[int]) -> Any:
    cell_type = cell.get("t", "n")
    style = cell.get("s")
    value_node = cell.find("m:v", _NS)
    if cell_type == "inlineStr":
        node = cell.find("m:is", _NS)
        return "".join(t.text or "" for t in node.iter(f"{{{_NS['m']}}}t")) if node is not None else ""
    if value_node is None or value_node.text is None:
        return ""
    raw = value_node.text
    if cell_type == "s":
        try:
            return shared[int(raw)]
        except (ValueError, IndexError):
            return ""
    if cell_type == "b":
        return raw == "1"
    if cell_type in {"str", "e"}:
        return raw
    if cell_type == "d":
        return raw
    try:
        number = float(raw)
    except ValueError:
        return raw
    if style is not None:
        try:
            if int(style) in date_styles:
                return _excel_serial_to_date(number)
            if int(style) in percent_styles:
                return _percent_value(number)
        except ValueError:
            pass
    if number.is_integer():
        return int(number)
    return number


def _merged_range(ref: str) -> tuple[int, int, int, int] | None:
    """A merged range such as D2:I2 as (first row, first column, last row, last column), 0-based."""

    parts = [_CELL_REF.fullmatch(part) for part in ref.split(":")]
    if len(parts) != 2 or not all(parts):
        return None
    (first_col, first_row), (last_col, last_row) = [(_column_index(part.group(0)), int(part.group(2)) - 1) for part in parts]
    return first_row, first_col, last_row, last_col


def _read_sheet(archive: zipfile.ZipFile, path: str, shared: list[str], date_styles: set[int], percent_styles: set[int], merged: list[tuple[int, int, int, int]] | None = None) -> list[list[Any]]:
    """Stream a worksheet row by row so the whole XML tree is never held in memory.

    Merged ranges near the top of the sheet (where the headers are) go into ``merged``.
    """

    rows: list[list[Any]] = []
    row_tag = f"{{{_NS['m']}}}row"
    cell_tag = f"{{{_NS['m']}}}c"
    merge_tag = f"{{{_NS['m']}}}mergeCell"
    values: list[Any] | None = None
    with open_entry(archive, path) as source:
        # Cells are read and released one at a time, so even a single row
        # holding millions of cells never builds a tree in memory.
        for event, element in ElementTree.iterparse(source, events=("start", "end")):
            if event == "start":
                if element.tag == row_tag:
                    values = []
                continue
            if element.tag == cell_tag:
                if values is not None:
                    index = _column_index(element.get("r", ""))
                    if index < MAX_SHEET_COLUMNS:
                        while len(values) < index:
                            values.append("")
                        values.append(_cell_value(element, shared, date_styles, percent_styles))
                element.clear()
                continue
            if element.tag == merge_tag:
                span = _merged_range(element.get("ref", ""))
                if merged is not None and span and span[0] < 40:
                    merged.append(span)
                element.clear()
                continue
            if element.tag == row_tag:
                rows.append(values if values is not None else [])
                values = None
                element.clear()
                if len(rows) > MAX_ROWS + 20:
                    break
    return rows


def _open_archive(content: bytes) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(BytesIO(content))
    except zipfile.BadZipFile as exc:
        raise ParserError("workbook is not a valid .xlsx archive") from exc
    try:
        check_archive_limits(archive)
    except ParserError:
        archive.close()
        raise
    return archive


def _parse_with_stdlib(file_name: str, content: bytes) -> ParseResult:
    with _open_archive(content) as archive:
        shared = _shared_strings(archive)
        date_styles, percent_styles = _number_styles(archive)
        result = ParseResult(file_name=file_name, file_kind=FileKind.XLSX, metadata={"engine": "stdlib"})
        for sheet_name, path in _sheet_paths(archive):
            merged: list[tuple[int, int, int, int]] = []
            try:
                rows = _read_sheet(archive, path, shared, date_styles, percent_styles, merged)
            except (KeyError, ElementTree.ParseError, zipfile.BadZipFile):
                result.warnings.append(f"sheet_unreadable:{sheet_name}")
                continue
            table = grid_to_table(rows, name=sheet_name, source_file=file_name, sheet=sheet_name, merged=merged)
            result.tables.append(table)
        result.page_count = len(result.tables)
    return result


def _openpyxl_value(cell: Any) -> Any:
    value = cell.value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            number_format = cell.number_format or ""
        except (IndexError, KeyError):  # a style or custom format the workbook never defines
            number_format = ""
        if _is_percent_format(number_format):
            return _percent_value(value)
    return value


def _parse_with_openpyxl(file_name: str, content: bytes) -> ParseResult:
    import openpyxl  # type: ignore[import-not-found]

    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    result = ParseResult(file_name=file_name, file_kind=FileKind.XLSX, metadata={"engine": "openpyxl"})
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            continue
        rows: list[list[Any]] = []
        for row in sheet.iter_rows():
            rows.append([_openpyxl_value(cell) for cell in row])
            if len(rows) > MAX_ROWS + 20:
                break
        result.tables.append(grid_to_table(rows, name=sheet.title, source_file=file_name, sheet=sheet.title))
    result.page_count = len(result.tables)
    return result


def parse_xlsx(file_name: str, content: bytes) -> ParseResult:
    # The inflated-size check runs before either engine touches an entry.
    _open_archive(content).close()
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return _parse_with_stdlib(file_name, content)
    try:
        return _parse_with_openpyxl(file_name, content)
    except Exception:  # noqa: BLE001 - fall back to the dependency-free reader
        return _parse_with_stdlib(file_name, content)


def parse_xls(file_name: str, content: bytes) -> ParseResult:
    try:
        import xlrd  # type: ignore[import-not-found]
    except ImportError as exc:
        raise ParserUnavailable("legacy .xls workbooks need the optional xlrd package; save the file as .xlsx or .csv") from exc
    book = xlrd.open_workbook(file_contents=content)
    result = ParseResult(file_name=file_name, file_kind=FileKind.XLS, metadata={"engine": "xlrd"})
    for sheet in book.sheets():
        rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(min(sheet.nrows, MAX_ROWS + 20))]
        result.tables.append(grid_to_table(rows, name=sheet.name, source_file=file_name, sheet=sheet.name))
    result.page_count = len(result.tables)
    return result


def tables_only(result: ParseResult) -> list[ParsedTable]:
    return [table for table in result.tables if table.headers]


__all__ = ["parse_xls", "parse_xlsx", "tables_only"]

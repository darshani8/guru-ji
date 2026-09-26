import importlib.util
import io
import json
import unittest
import zipfile

from app.ingestion import parsers
from app.ingestion.detector import detect_file_kind
from app.ingestion.models import FileKind, ParserError, ParserUnavailable
from app.ingestion.parsers import archive as archive_guard
from app.ingestion.parsers import excel_parser
from app.ingestion.parsers import pdf_parser
from app.ingestion.parsers.csv_parser import parse_csv
from app.ingestion.parsers.docx_parser import parse_docx
from app.ingestion.parsers.excel_parser import parse_xlsx
from app.ingestion.parsers.image_parser import parse_image
from app.ingestion.parsers.json_parser import parse_json
from app.ingestion.parsers.ocr import OcrResult, OcrUnavailable, parse_textract_blocks
from app.ingestion.parsers.tabular import detect_header_row, grid_to_table
from app.ingestion.parsers.text_layout import key_value_fields, text_to_table
from app.ingestion.registry import ParserRegistry
from app.ingestion.service import JOB_IMPORTED, IngestionService
from app.institution_data.store import InstitutionDataStore
from app.storage.object_store import InMemoryObjectStore

OPENPYXL_AVAILABLE = importlib.util.find_spec("openpyxl") is not None
# Cell styles build_xlsx writes; a cell given as (style, number) uses one of them.
PERCENT, PERCENT_2DP, CUSTOM_PERCENT, ESCAPED_PERCENT_SIGN, QUOTED_PERCENT_SIGN = 2, 3, 4, 5, 6


def build_xlsx(sheets: dict[str, list[list[object]]], merges: dict[str, list[str]] | None = None) -> bytes:
    """A minimal workbook; a None cell is left out, and ``merges`` lists merged ranges (such as "C2:E2") per sheet."""

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        overrides = "".join(f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>' for i in range(1, len(sheets) + 1))
        archive.writestr("[Content_Types].xml", f'<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>{overrides}</Types>')
        archive.writestr("_rels/.rels", '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        names = "".join(f'<sheet name="{name}" sheetId="{i}" r:id="rId{i}"/>' for i, name in enumerate(sheets, start=1))
        archive.writestr("xl/workbook.xml", f'<?xml version="1.0"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>{names}</sheets></workbook>')
        rels = "".join(f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>' for i in range(1, len(sheets) + 1))
        archive.writestr("xl/_rels/workbook.xml.rels", f'<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{rels}</Relationships>')
        strings = ["shared one"]
        archive.writestr("xl/sharedStrings.xml", '<?xml version="1.0"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>shared one</t></si></sst>')
        # Styles 2-4 show percentages (built-in 0% and 0.00%, custom [Blue]0.0%); 5 and 6 only print a % sign.
        formats = '<numFmts count="3"><numFmt numFmtId="164" formatCode="[Blue]0.0%"/><numFmt numFmtId="165" formatCode="0\\%"/><numFmt numFmtId="166" formatCode="0&quot;%&quot;"/></numFmts>'
        styles = "".join(f'<xf numFmtId="{fmt}"/>' for fmt in (0, 14, 9, 10, 164, 165, 166))
        archive.writestr("xl/styles.xml", f'<?xml version="1.0"?><styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">{formats}<cellXfs count="7">{styles}</cellXfs></styleSheet>')
        for index, (sheet_name, rows) in enumerate(sheets.items(), start=1):
            xml = ['<?xml version="1.0"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>']
            for r, row in enumerate(rows, start=1):
                cells = []
                for c, value in enumerate(row):
                    ref = f"{chr(65 + c)}{r}"
                    if value is None:
                        continue
                    if value == "__shared__":
                        cells.append(f'<c r="{ref}" t="s"><v>0</v></c>')
                    elif value == "__date__":
                        cells.append(f'<c r="{ref}" s="1"><v>45000</v></c>')
                    elif isinstance(value, tuple):
                        style, number = value
                        cells.append(f'<c r="{ref}" s="{style}"><v>{number}</v></c>')
                    elif isinstance(value, bool):
                        cells.append(f'<c r="{ref}" t="b"><v>{int(value)}</v></c>')
                    elif isinstance(value, (int, float)):
                        cells.append(f'<c r="{ref}"><v>{value}</v></c>')
                    else:
                        cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{value}</t></is></c>')
                xml.append(f'<row r="{r}">' + "".join(cells) + "</row>")
            xml.append("</sheetData>")
            ranges = (merges or {}).get(sheet_name, [])
            if ranges:
                xml.append(f'<mergeCells count="{len(ranges)}">' + "".join(f'<mergeCell ref="{ref}"/>' for ref in ranges) + "</mergeCells>")
            xml.append("</worksheet>")
            archive.writestr(f"xl/worksheets/sheet{index}.xml", "".join(xml))
        del strings
    return buffer.getvalue()


def build_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    body = "".join(f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>" for text in paragraphs)
    if table:
        rows = "".join("<w:tr>" + "".join(f"<w:tc><w:p><w:r><w:t>{cell}</w:t></w:r></w:p></w:tc>" for cell in row) + "</w:tr>" for row in table)
        body += f"<w:tbl>{rows}</w:tbl>"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="application/xml"/></Types>')
        archive.writestr("word/document.xml", f'<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>{body}</w:body></w:document>')
    return buffer.getvalue()


def rewrite_entry(content: bytes, name: str, transform) -> bytes:
    """Rebuild a zip archive (deflated) with one entry's bytes transformed."""

    source = zipfile.ZipFile(io.BytesIO(content))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as out:
        for info in source.infolist():
            data = source.read(info.filename)
            out.writestr(info.filename, transform(data) if info.filename == name else data)
    return buffer.getvalue()


class _FixtureOcr:
    engine_name = "fixture"

    def __init__(self, text: str, confidence: float = 0.9):
        self.text = text
        self.confidence = confidence

    def recognize(self, image_bytes: bytes, *, content_type: str = "image/png") -> OcrResult:
        return OcrResult(text=self.text, confidence=self.confidence, engine=self.engine_name)


class ParserTests(unittest.TestCase):
    def test_detection_prefers_magic_bytes_over_extension(self):
        self.assertEqual(detect_file_kind("scan.xlsx", b"%PDF-1.4 ..."), FileKind.PDF)
        self.assertEqual(detect_file_kind("photo.bin", b"\x89PNG\r\n"), FileKind.IMAGE)
        self.assertEqual(detect_file_kind("students.csv", b"a,b\n1,2"), FileKind.CSV)
        self.assertEqual(detect_file_kind("data", b"a\tb\n1\t2", "text/tab-separated-values"), FileKind.TSV)
        self.assertEqual(detect_file_kind("x.docx", build_docx(["hi"])), FileKind.DOCX)
        self.assertEqual(detect_file_kind("x.xlsx", build_xlsx({"S": [["a"]]})), FileKind.XLSX)
        self.assertEqual(detect_file_kind("blob", b"\x00\x01\x02"), FileKind.UNKNOWN)

    def test_csv_header_detection_skips_titles_and_dedupes_headers(self):
        content = "ABC College\nStudent list 2026\n\nName,USN,Sem,Name\nRavi,1MS23MBA001,1,x\n,,,\nAsha,1MS23MBA002,2,y\n".encode("utf-8-sig")
        result = parse_csv("students.csv", content)
        table = result.tables[0]
        self.assertEqual(table.headers, ("Name", "USN", "Sem", "Name (2)"))
        self.assertEqual(table.row_count, 2)
        self.assertEqual(table.records[0].locator, "csv;row=5")
        self.assertIn("skipped_3_leading_rows", table.warnings)
        self.assertIsNone(detect_header_row([["1", "2"], ["3", "4"]]))

    def test_xlsx_reader_handles_shared_strings_dates_and_multiple_sheets(self):
        content = build_xlsx({"MBA": [["Name", "Student ID", "Joined", "Active"], ["__shared__", "MBA001", "__date__", True], ["Asha", 42, "2026-01-05", False]], "Notes": [["only a title"]]})
        result = parse_xlsx("students.xlsx", content)
        self.assertEqual(result.metadata["engine"], "stdlib")
        table = next(item for item in result.tables if item.name == "MBA")
        self.assertEqual(table.headers, ("Name", "Student ID", "Joined", "Active"))
        self.assertEqual(table.records[0].fields["Name"], "shared one")
        self.assertEqual(table.records[0].fields["Joined"], "2023-03-15")
        self.assertIs(table.records[0].fields["Active"], True)
        self.assertEqual(table.records[1].fields["Student ID"], 42)

    def test_xlsx_reader_follows_sheet_paths_from_the_package_root(self):
        # Some writers point at /xl/worksheets/sheet1.xml instead of worksheets/sheet1.xml.
        content = rewrite_entry(build_xlsx({"S": [["USN", "Name"], ["1MS23MBA001", "Ravi"]]}), "xl/_rels/workbook.xml.rels", lambda data: data.replace(b'Target="worksheets/', b'Target="/xl/worksheets/'))
        result = excel_parser._parse_with_stdlib("students.xlsx", content)
        self.assertEqual(result.warnings, [])
        self.assertEqual(result.tables[0].headers, ("USN", "Name"))
        self.assertEqual(result.tables[0].row_count, 1)

    def test_percent_formatted_cells_read_as_the_percentage_shown(self):
        # Excel stores 85% as 0.85 and 100% as 1; the reader returns what the cell shows.
        content = build_xlsx({"Att": [
            ["USN", "Built-in", "Two places", "Custom", "Escaped sign", "Quoted sign", "Plain", "Joined"],
            ["1MS23MBA001", (PERCENT, 0), (PERCENT_2DP, 0.85), (CUSTOM_PERCENT, 0.125), (ESCAPED_PERCENT_SIGN, 85), (QUOTED_PERCENT_SIGN, 85), 1, "__date__"],
            ["1MS23MBA002", (PERCENT, 1), (PERCENT_2DP, 0.07), (CUSTOM_PERCENT, 1.5), (ESCAPED_PERCENT_SIGN, 1), (QUOTED_PERCENT_SIGN, 1), 0.85, "__date__"],
            ["1MS23MBA003", (PERCENT, 0.85), (PERCENT_2DP, 0.005), (CUSTOM_PERCENT, 0.01), (ESCAPED_PERCENT_SIGN, 0), (QUOTED_PERCENT_SIGN, 0), 0, "__date__"],
        ]})
        result = excel_parser._parse_with_stdlib("attendance.xlsx", content)
        rows = [record.fields for record in result.tables[0].records]
        self.assertEqual([row["Built-in"] for row in rows], [0, 100, 85])
        self.assertEqual([row["Two places"] for row in rows], [85, 7, "0.5%"], "no float noise, and a value under 1% keeps its sign so it is not read as a fraction")
        self.assertEqual([row["Custom"] for row in rows], [12.5, 150, 1])
        self.assertEqual([row["Escaped sign"] for row in rows], [85, 1, 0], "a format that only prints a % sign does not scale")
        self.assertEqual([row["Quoted sign"] for row in rows], [85, 1, 0])
        self.assertEqual([row["Plain"] for row in rows], [1, 0.85, 0])
        self.assertEqual([row["Joined"] for row in rows], ["2023-03-15"] * 3)
        self.assertIs(type(rows[2]["Custom"]), int, "1% stays the whole number 1; a float 1.0 would be read later as the fraction 100%")

    def test_a_number_format_the_workbook_never_defines_reads_the_value_as_stored(self):
        # openpyxl's read-only cells look the format up by index; a style or
        # custom format id missing from the workbook raises IndexError there.
        class BrokenStyleCell:
            value = 0.85

            @property
            def number_format(self):
                raise IndexError("list index out of range")

        self.assertEqual(excel_parser._openpyxl_value(BrokenStyleCell()), 0.85)

    @unittest.skipUnless(OPENPYXL_AVAILABLE, "openpyxl is not installed")
    def test_openpyxl_reads_percent_formatted_cells_as_the_percentage_shown(self):
        from datetime import date

        import openpyxl

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["USN", "Attendance %", "Custom", "Escaped sign", "Plain", "Joined", "Flag"])
        for row in (["1MS23MBA001", 0, 0.125, 85, 1, date(2026, 1, 5), True], ["1MS23MBA002", 1, 1.5, 1, 0.85, date(2026, 1, 5), False], ["1MS23MBA003", 0.85, 0.005, 0, 0, date(2026, 1, 5), True]):
            sheet.append(row)
        for number_format, column in (("0%", "B"), ("[Blue]0.0%", "C"), ("0\\%", "D"), ("0%", "G")):
            for cell in sheet[column][1:]:
                cell.number_format = number_format
        buffer = io.BytesIO()
        workbook.save(buffer)
        result = excel_parser._parse_with_openpyxl("attendance.xlsx", buffer.getvalue())
        rows = [record.fields for record in result.tables[0].records]
        self.assertEqual([row["Attendance %"] for row in rows], [0, 100, 85])
        self.assertEqual([row["Custom"] for row in rows], [12.5, 150, "0.5%"])
        self.assertEqual([row["Escaped sign"] for row in rows], [85, 1, 0])
        self.assertEqual([row["Plain"] for row in rows], [1, 0.85, 0])
        self.assertEqual([row["Joined"] for row in rows], ["2026-01-05T00:00:00"] * 3)
        self.assertEqual([row["Flag"] for row in rows], [True, False, True])

    def test_docx_parser_returns_text_and_tables(self):
        result = parse_docx("staff.docx", build_docx(["Faculty list", "Department of MBA"], [["Name", "Designation"], ["Dr Meena", "Professor"], ["Mr Rao", "Assistant Professor"]]))
        self.assertEqual(result.tables[0].headers, ("Name", "Designation"))
        self.assertEqual(result.tables[0].records[1].fields["Name"], "Mr Rao")
        self.assertIn("Department of MBA", result.full_text())

    def test_json_parser_finds_row_arrays(self):
        payload = {"meta": {"x": 1}, "students": [{"name": "A", "usn": "1"}, {"name": "B", "usn": "2", "extra": True}]}
        result = parse_json("export.json", json.dumps(payload).encode())
        self.assertEqual(result.tables[0].headers, ("name", "usn", "extra"))
        self.assertEqual(result.tables[0].row_count, 2)

    def test_text_layout_detects_tables_and_form_fields(self):
        text = "USN          Name         Attendance\n1MS23MBA001  Ravi Kumar   72%\n1MS23MBA002  Asha Rao     91%\n1MS23MBA003  Kiran        65%\n"
        table = text_to_table(text, name="layout", source_file="scan.png", page=1, ocr=True, ocr_confidence=0.9)
        self.assertIsNotNone(table)
        self.assertEqual(table.headers, ("USN", "Name", "Attendance"))
        self.assertEqual(table.records[0].fields["Attendance"], "72%")
        fields = key_value_fields("Name: Ravi Kumar\nDate of Birth: 12/05/2003\nProgram - MBA\nSignature")
        self.assertEqual(fields["Name"], "Ravi Kumar")
        self.assertEqual(fields["Program"], "MBA")

    def test_image_parser_uses_ocr_engine_and_flags_low_confidence(self):
        result = parse_image("form.png", b"\x89PNG", ocr_engine=_FixtureOcr("Name: Ravi\nUSN: 1MS23MBA001\nProgram: MBA\nSem: 1", confidence=0.6))
        self.assertEqual(result.tables[0].name, "form_fields")
        self.assertEqual(result.tables[0].records[0].fields["USN"], "1MS23MBA001")
        self.assertTrue(result.tables[0].records[0].ocr)
        self.assertIn("low_ocr_confidence", result.warnings)
        with self.assertRaises(ParserUnavailable):
            parse_image("form.png", b"\x89PNG")

    def test_textract_blocks_become_tables(self):
        blocks = [
            {"Id": "l1", "BlockType": "LINE", "Text": "Attendance register", "Confidence": 99.0},
            {"Id": "t1", "BlockType": "TABLE", "Relationships": [{"Type": "CHILD", "Ids": ["c1", "c2", "c3", "c4"]}]},
            {"Id": "c1", "BlockType": "CELL", "RowIndex": 1, "ColumnIndex": 1, "Relationships": [{"Type": "CHILD", "Ids": ["w1"]}]},
            {"Id": "c2", "BlockType": "CELL", "RowIndex": 1, "ColumnIndex": 2, "Relationships": [{"Type": "CHILD", "Ids": ["w2"]}]},
            {"Id": "c3", "BlockType": "CELL", "RowIndex": 2, "ColumnIndex": 1, "Relationships": [{"Type": "CHILD", "Ids": ["w3"]}]},
            {"Id": "c4", "BlockType": "CELL", "RowIndex": 2, "ColumnIndex": 2, "Relationships": [{"Type": "CHILD", "Ids": ["w4"]}]},
            {"Id": "w1", "BlockType": "WORD", "Text": "USN"}, {"Id": "w2", "BlockType": "WORD", "Text": "Present"},
            {"Id": "w3", "BlockType": "WORD", "Text": "MBA001"}, {"Id": "w4", "BlockType": "WORD", "Text": "12"},
        ]
        result = parse_textract_blocks(blocks)
        self.assertEqual(result.tables[0], (("USN", "Present"), ("MBA001", "12")))
        self.assertEqual(result.confidence, 0.99)

    def test_pdf_requires_an_engine_and_decides_ocr_page_by_page(self):
        with self.assertRaises(ParserUnavailable):
            pdf_parser.parse_pdf("x.pdf", b"%PDF-1.4\n")
        original = pdf_parser.extract_pdf_pages
        pdf_parser.extract_pdf_pages = lambda content: ([
            {"number": 1, "text": "Fee policy\nTuition is payable in two instalments each semester, with penalties after the due date.", "tables": [[["USN", "Paid"], ["MBA001", "5000"]]], "image": None},
            {"number": 2, "text": "", "tables": [], "image": b"\x89PNG"},
            {"number": 3, "text": "", "tables": [], "image": None},
        ], "fake")
        try:
            result = pdf_parser.parse_pdf("fees.pdf", b"%PDF-1.4\n", ocr_engine=_FixtureOcr("USN   Paid\nMBA002  4000\nMBA003  3000\nMBA004  2500"))
        finally:
            pdf_parser.extract_pdf_pages = original
        self.assertEqual(result.metadata["ocr_pages"], 1)
        self.assertEqual(result.metadata["text_pages"], 1)
        self.assertIn("page_3_needs_ocr", result.warnings)
        self.assertEqual(result.tables[0].records[0].fields["USN"], "MBA001")
        ocr_table = next(table for table in result.tables if table.ocr)
        self.assertEqual(ocr_table.records[0].fields["Paid"], "4000")
        self.assertEqual(len(result.texts), 2)

    def test_registry_rejects_empty_and_oversized_files(self):
        registry = ParserRegistry(max_bytes=10)
        with self.assertRaises(ValueError):
            registry.parse("x.csv", b"")
        with self.assertRaises(ValueError):
            registry.parse("x.csv", b"a,b\n" * 10)
        self.assertIn(FileKind.IMAGE, registry.supported_kinds())
        self.assertTrue(parsers.__doc__)

    def test_parsers_raise_parser_error_for_their_own_failure_modes(self):
        # An unbalanced quote swallows the rest of the file into one field (csv.Error).
        with self.assertRaises(ParserError):
            parse_csv("bad.csv", b"a,b\n\"" + b"x" * 200_000 + b"\n1,2\n")
        # A zip with xl/ entries but no workbook index (KeyError before).
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as bad:
            bad.writestr("xl/other.xml", "<x/>")
        with self.assertRaises(ParserError):
            parse_xlsx("bad.xlsx", buffer.getvalue())
        # Malformed workbook parts are reported, not raised as ElementTree.ParseError.
        broken = rewrite_entry(build_xlsx({"S": [["a", "b"], [1, 2]]}), "xl/sharedStrings.xml", lambda data: b"<sst><si><t>unterminated")
        with self.assertRaises(ParserError):
            parse_xlsx("broken.xlsx", broken)
        broken_sheet = rewrite_entry(build_xlsx({"S": [["a", "b"], [1, 2]]}), "xl/worksheets/sheet1.xml", lambda data: data[: len(data) // 2])
        self.assertIn("sheet_unreadable:S", parse_xlsx("broken.xlsx", broken_sheet).warnings)
        with self.assertRaises(ParserError):
            parse_docx("bad.docx", rewrite_entry(build_docx(["hi"]), "word/document.xml", lambda data: data[:20]))
        # Deeply nested JSON (RecursionError before).
        with self.assertRaises(ParserError):
            parse_json("deep.json", b"[" * 100_000 + b"]" * 100_000)

    def test_zip_bombs_are_rejected_before_inflating(self):
        base = build_xlsx({"S": [["a", "b"], [1, 2]]})
        # 1.5 MB of padding deflates to a few KB: ratio far above 100:1.
        pad = lambda data: data.replace(b"</sheetData>", b"<!--" + b" " * 1_500_000 + b"--></sheetData>")
        bomb = rewrite_entry(base, "xl/worksheets/sheet1.xml", pad)
        self.assertLess(len(bomb), 20_000)
        with self.assertRaises(ParserError) as raised:
            parse_xlsx("bomb.xlsx", bomb)
        self.assertIn("compression ratio", str(raised.exception))
        with self.assertRaises(ParserError):
            parse_xlsx("bomb.xlsx", rewrite_entry(base, "xl/sharedStrings.xml", lambda data: data.replace(b"</sst>", b"<!--" + b" " * 1_500_000 + b"--></sst>")))
        doc_bomb = rewrite_entry(build_docx(["hi"]), "word/document.xml", lambda data: data.replace(b"</w:body>", b"<!--" + b" " * 1_500_000 + b"--></w:body>"))
        with self.assertRaises(ParserError):
            parse_docx("bomb.docx", doc_bomb)
        # Below the ratio floor the total inflated size cap still applies.
        small = rewrite_entry(base, "xl/worksheets/sheet1.xml", lambda data: data.replace(b"</sheetData>", b"<!--" + b"y" * 300_000 + b"--></sheetData>"))
        original = archive_guard.MAX_INFLATED_BYTES
        archive_guard.MAX_INFLATED_BYTES = 200_000
        try:
            with self.assertRaises(ParserError):
                parse_xlsx("large.xlsx", small)
            with self.assertRaises(ParserError):
                parse_docx("large.docx", rewrite_entry(build_docx(["hi"]), "word/document.xml", lambda data: data.replace(b"</w:body>", b"<!--" + b"y" * 300_000 + b"--></w:body>")))
        finally:
            archive_guard.MAX_INFLATED_BYTES = original
        self.assertEqual(parse_xlsx("ok.xlsx", small).tables[0].row_count, 1)
        # The byte budget is enforced on delivered bytes, not only on declared sizes.
        reader = archive_guard.BoundedReader(io.BytesIO(b"z" * 1000), budget=500)
        with self.assertRaises(ParserError):
            while reader.read(200):
                pass

    def test_json_parser_caps_distinct_keys_and_keeps_only_present_keys(self):
        result = parse_json("sparse.json", json.dumps([{"a": 1}, {"b": 2, "a": 3}]).encode())
        self.assertEqual(result.tables[0].headers, ("a", "b"))
        self.assertEqual([record.fields for record in result.tables[0].records], [{"a": 1}, {"b": 2, "a": 3}])
        wide = [{f"k{i}": i} for i in range(201)]
        with self.assertRaises(ParserError):
            parse_json("wide.json", json.dumps(wide).encode())
        self.assertEqual(len(parse_json("ok.json", json.dumps(wide[:200]).encode()).tables[0].headers), 200)

    def test_disabled_ocr_engine_reports_configuration(self):
        from app.ingestion.parsers.ocr import DisabledOcrEngine

        with self.assertRaises(OcrUnavailable):
            DisabledOcrEngine().recognize(b"x")


class PercentImportTests(unittest.IsolatedAsyncioTestCase):
    async def test_percent_formatted_attendance_imports_as_the_percentage_shown(self):
        store = InstitutionDataStore(":memory:")
        service = IngestionService(store=store, objects=InMemoryObjectStore(), parsers=ParserRegistry())
        content = build_xlsx({"Aug": [
            ["USN", "Subject Code", "Month", "Attendance %"],
            ["1MS23MBA001", "MBA101", "Aug", (PERCENT, 1)],
            ["1MS23MBA002", "MBA101", "Aug", (PERCENT, 0.85)],
            ["1MS23MBA003", "MBA101", "Aug", (PERCENT_2DP, 0.005)],
            ["1MS23MBA004", "MBA101", "Aug", (PERCENT, 0)],
            ["1MS23MBA005", "MBA101", "Aug", (PERCENT, 0.01)],
        ]})
        job = service.upload("college_a", "staff-1", file_name="attendance.xlsx", content=content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        job = await service.process("college_a", job["job_id"])
        self.assertEqual((job["entity"], job["status"]), ("attendance", JOB_IMPORTED))
        stored = {row["student_id"]: row["attendance_percent"] for row in store.query_records("college_a", "attendance")}
        self.assertEqual(stored, {"1MS23MBA001": 100, "1MS23MBA002": 85, "1MS23MBA003": 0.5, "1MS23MBA004": 0, "1MS23MBA005": 1})


class HeaderRowTests(unittest.TestCase):
    """The first data row must never be read as the headers, and no data row or column may be lost."""

    def test_a_blank_header_cell_does_not_make_the_first_student_the_header(self):
        cases = {
            "unlabelled trailing column": ("USN,Student Name,Program,Semester,\n1MS23MBA001,Ravi Kumar,MBA,1,ok\n1MS23MBA002,Asha Rao,MBA,1,\n", ("USN", "Student Name", "Program", "Semester", "column_5")),
            "blank cell in the middle": ("USN,Student Name,,Program,Semester\n1MS23MBA001,Ravi Kumar,A,MBA,1\n1MS23MBA002,Asha Rao,B,MBA,1\n", ("USN", "Student Name", "column_3", "Program", "Semester")),
            "blank first cell over serial numbers": (",USN,Student Name,Semester,Mobile\n1,1MS23MBA001,Ravi Kumar,1,9876543210\n2,1MS23MBA002,Asha Rao,1,9876543211\n", ("column_1", "USN", "Student Name", "Semester", "Mobile")),
            "blank first cell over serial numbers, text columns": (",Name,Program,Section\n1,Ravi Kumar,MBA,A\n2,Asha Rao,MCA,B\n", ("column_1", "Name", "Program", "Section")),
            "text columns only": ("Name,Program,Section,\nRavi Kumar,MBA,A,ok\nAsha Rao,MCA,B,\n", ("Name", "Program", "Section", "column_4")),
        }
        for label, (content, headers) in cases.items():
            with self.subTest(label):
                table = parse_csv("students.csv", content.encode()).tables[0]
                self.assertEqual(table.headers, headers)
                self.assertEqual(table.row_count, 2)
                self.assertNotIn("skipped_1_leading_rows", table.warnings)
        # A key/value line above, and a blank row between the header and the students.
        table = parse_csv("students.csv", b"Subject,Marketing\nUSN,Name,IA1,\n\n1MS23MBA001,Ravi,20,ok\n1MS23MBA002,Asha,22,\n").tables[0]
        self.assertEqual(table.headers, ("USN", "Name", "IA1", "column_4"))
        self.assertEqual([record.fields["USN"] for record in table.records], ["1MS23MBA001", "1MS23MBA002"])

    def test_a_short_header_row_keeps_the_unlabelled_column(self):
        # A spreadsheet row ends at its last filled cell, so the header row is one cell short.
        workbook = build_xlsx({"Students": [["USN", "Student Name", "Program", "Semester"], ["1MS23MBA001", "Ravi Kumar", "MBA", 1, "ok"], ["1MS23MBA002", "Asha Rao", "MBA", 1, "late"]]})
        table = parse_xlsx("students.xlsx", workbook).tables[0]
        self.assertEqual(table.headers, ("USN", "Student Name", "Program", "Semester", "column_5"))
        self.assertEqual([record.fields["USN"] for record in table.records], ["1MS23MBA001", "1MS23MBA002"])
        self.assertEqual([record.fields["column_5"] for record in table.records], ["ok", "late"])
        # Trailing blank cells in a data row do not add columns.
        table = grid_to_table([["USN", "Name"], ["1MS23MBA001", "Ravi", "", None]], name="t", source_file="t.csv")
        self.assertEqual(table.headers, ("USN", "Name"))

    def test_year_and_day_headers_are_headers(self):
        table = parse_csv("marks.csv", b"USN,Name,2023,2024\n1MS23MBA001,Ravi,10,20\n1MS23MBA002,Asha,11,21\n").tables[0]
        self.assertEqual(table.headers, ("USN", "Name", "2023", "2024"))
        self.assertEqual(table.records[0].fields["2024"], "20")
        workbook = build_xlsx({
            "Marks": [["USN", "Name", 2023, 2024], ["1MS23MBA001", "Ravi", 10, 20], ["1MS23MBA002", "Asha", 11, 21]],
            "Latest": [["USN", "Name", 2023, 2024], ["1MS23MBA001", "Ravi", 10, 20]],
            "Intake": [["Programme", 2021, 2022, 2023], ["MBA", 120, 118, 125], ["MCA", 60, 58, 61]],
        })
        tables = {table.name: table for table in parse_xlsx("years.xlsx", workbook).tables}
        self.assertEqual(tables["Marks"].headers, ("USN", "Name", "2023", "2024"))
        self.assertEqual(tables["Marks"].records[0].fields["USN"], "1MS23MBA001")
        self.assertEqual(tables["Latest"].headers, ("USN", "Name", "2023", "2024"))
        self.assertEqual(tables["Intake"].headers, ("Programme", "2021", "2022", "2023"))
        self.assertEqual(tables["Intake"].row_count, 2)
        days = ",".join(str(day) for day in range(1, 32))
        students = "".join(f"1MS23MBA00{i},Student {i}," + ",".join("A" if (day + i) % 6 == 0 else "P" for day in range(31)) + ",26\n" for i in range(1, 4))
        table = parse_csv("attendance.csv", f"USN,Name,{days},Total\n{students}".encode()).tables[0]
        self.assertEqual(table.headers[:4], ("USN", "Name", "1", "2"))
        self.assertEqual(table.headers[-2:], ("31", "Total"))
        self.assertEqual(table.row_count, 3)

    def test_xls_float_numbers_in_the_header_row(self):
        # xlrd reads every number as a float: 2023 arrives as 2023.0 and day 1 as 1.0.
        grids = {
            "years": ([["USN", "Name", 2023.0, 2024.0], ["1MS23MBA001", "Ravi", 10.0, 20.0], ["1MS23MBA002", "Asha", 11.0, 21.0]], ("USN", "Name", "2023", "2024")),
            "descending years": ([["USN", "Name", 2024.0, 2023.0], ["1MS23MBA001", "Ravi", 10.0, 20.0], ["1MS23MBA002", "Asha", 11.0, 21.0]], ("USN", "Name", "2024", "2023")),
            "intake": ([["Programme", 2021.0, 2022.0, 2023.0], ["MBA", 120.0, 118.0, 125.0], ["MCA", 60.0, 58.0, 61.0]], ("Programme", "2021", "2022", "2023")),
            "days": ([["USN", "Name", *[float(day) for day in range(1, 8)], "Total"], ["1MS23MBA001", "Ravi", *"PPAPPPA", 5.0], ["1MS23MBA002", "Asha", *"PPPPPAP", 6.0]], ("USN", "Name", "1", "2", "3", "4", "5", "6", "7", "Total")),
        }
        for label, (grid, headers) in grids.items():
            with self.subTest(label):
                table = grid_to_table(grid, name="Sheet1", source_file="book.xls", sheet="Sheet1")
                self.assertEqual(table.headers, headers)
                self.assertEqual(table.row_count, 2)
                self.assertEqual(table.warnings, [])
                # Data values are left as the reader gave them.
                self.assertEqual(table.records[0].fields[headers[-1]], grid[1][-1])

    def test_every_data_row_is_kept_when_the_header_repeats(self):
        sections = "Section A,,,\nUSN,Name,Program,Semester\n1MS23MBA001,Ravi Kumar,MBA,3\n1MS23MBA002,Asha Rao,MBA,3\n1MS23MBA003,Kiran S,MBA,3\nSection B,,,\nUSN,Name,Program,Semester\n1MS23MBA004,Meena P,MBA,3\n1MS23MBA005,Arjun K,MBA,3\n"
        table = parse_csv("sections.csv", sections.encode()).tables[0]
        self.assertEqual(table.headers, ("USN", "Name", "Program", "Semester"))
        self.assertEqual(table.warnings, ["skipped_1_leading_rows"])
        self.assertTrue({f"1MS23MBA00{i}" for i in range(1, 6)} <= {record.fields["USN"] for record in table.records})
        workbook = build_xlsx({"Students": [["Section A"], ["USN", "Name", "Program", "Semester"], ["1MS23MBA001", "Ravi Kumar", "MBA", 3], ["1MS23MBA002", "Asha Rao", "MBA", 3], ["1MS23MBA003", "Kiran S", "MBA", 3], ["Section B"], ["USN", "Name", "Program", "Semester"], ["1MS23MBA004", "Meena P", "MBA", 3]]})
        table = parse_xlsx("sections.xlsx", workbook).tables[0]
        self.assertEqual(table.headers, ("USN", "Name", "Program", "Semester"))
        self.assertTrue({f"1MS23MBA00{i}" for i in range(1, 5)} <= {record.fields["USN"] for record in table.records})
        grids = {
            "header repeated after 2 rows (all text)": [["Name", "Program", "Section"], ["Ravi", "MBA", "A"], ["Asha", "MCA", "B"], ["Name", "Program", "Section"], ["Kiran", "MBA", "A"], ["Meena", "MCA", "B"]],
            "header repeated after 3 rows": [["USN", "Name", "Program", "Section"], ["1MS23MBA001", "Ravi", "MBA", "A"], ["1MS23MBA002", "Asha", "MBA", "A"], ["1MS23MBA003", "Kiran", "MBA", "A"], ["USN", "Name", "Program", "Section"], ["1MS23MBA004", "Meena", "MBA", "B"]],
            "header with a blank cell repeated": [["Name", "Program", ""], ["Ravi", "MBA", "ok"], ["Name", "Program", ""], ["Kiran", "MCA", "ok"]],
        }
        for label, grid in grids.items():
            with self.subTest(label):
                self.assertEqual(detect_header_row(grid), 0)

    def test_tables_the_width_score_already_read_correctly_are_unchanged(self):
        faculty = [["Ravi Kumar", "Professor", "Marketing"], ["Asha Rao", "Associate Professor", "Finance"], ["Kiran S", "Assistant Professor", "HR"]]
        students = [["Ravi Kumar", "MBA", "A"], ["Asha Rao", "MCA", "B"], ["Kiran S", "BBA", "C"]]
        timetable = [["Mon", "18MBA11", "18MBA12", "18MBA13", "18MBA14"], ["Tue", "18MBA12", "18MBA13", "18MBA14", "18MBA11"], ["Wed", "18MBA13", "18MBA14", "18MBA11", "18MBA12"]]
        grids = {
            "group header line above a faculty header": ([["Faculty Details", "", "Department Info"], ["Name", "Designation", "Department"], *faculty], 1),
            "two-cell title above a faculty header": ([["ABC Institute of Management", "Bengaluru"], ["Name", "Designation", "Department"], *faculty], 1),
            "key/value lines above a text header": ([["Department", "MBA"], ["Semester", "III"], ["Name", "Program", "Section"], *students], 2),
            "key/value line with a colon above a text header": ([["Department:", "Management Studies"], ["Name", "Program", "Section"], *students], 1),
            "class and faculty lines above a text header": ([["Class", "MBA III Sem"], ["Faculty", "Dr. Rao"], ["Name", "Program", "Section"], *students], 2),
            "key/value line above a header with a blank cell": ([["Class:", "Marketing"], ["18MBA11", "CO1", "", "City"], ["O", "5", "Finance", "Bengaluru"]], 1),
            "timetable": ([["Day", "P1", "P2", "P3", "P4"], *timetable], 0),
            "timetable under a two-cell title": ([["Timetable", "Semester III"], ["Day", "P1", "P2", "P3", "P4"], *timetable], 1),
            "two-line header of equal width": ([["USN", "Name", "Marks", "", ""], ["", "", "Maths", "Physics", "Chem"], ["1MS23MBA001", "Ravi", "50", "60", "70"], ["1MS23MBA002", "Asha", "55", "65", "75"]], 0),
            "subject codes over grades": ([["USN", "Name", "18MBA11", "18MBA12", "18MBA13"], ["1MS23MBA001", "Ravi", "A", "B+", "O"], ["1MS23MBA002", "Asha", "B", "A+", "A"]], 0),
            "answers that happen to count up": ([["Name", "Q1", "Q2", "Q3", ""], ["Ravi", "1", "2", "3", "good"], ["Asha", "4", "5", "3", "ok"]], 0),
            "a title and a date above the header": ([["ABC College Student List", "", "Date: 01-01-2025"], ["Name", "Program", "Section"], ["Ravi", "MBA", "A"], ["Asha", "MCA", "B"]], 1),
            "college details above the header": ([["Name of the Institution", "ABC College", "", "Academic Year", "2024-25", ""], ["USN", "Name", "Program", "Section", "Email", "Phone"], ["1MS23MBA001", "Ravi", "MBA", "A", "ravi@abc.edu", "9876543210"], ["1MS23MBA002", "Asha", "MBA", "B", "asha@abc.edu", "9876543211"]], 1),
            "batch years as data": ([["Name", "From", "To"], ["Ravi", "2023", "2024"], ["Asha", "2023", "2024"]], 0),
            "a header word used as a value": ([["Name", "Status", "Remarks"], ["Ravi", "Active", "Status pending"], ["Asha", "Active", "Remarks"]], 0),
            "titles above the header": ([["ABC College"], ["Student list 2025"], ["USN", "Name", "Program"], ["1MS23MBA001", "Ravi", "MBA"]], 2),
        }
        for label, (grid, expected) in grids.items():
            with self.subTest(label):
                self.assertEqual(detect_header_row(grid), expected)
        table = parse_csv("faculty.csv", b"Department:,Management Studies\nName,Designation,Qualification\nRavi Kumar,Professor,PhD\nAsha Rao,Associate Professor,MBA\nKiran S,Assistant Professor,M.Com\n").tables[0]
        self.assertEqual(table.headers, ("Name", "Designation", "Qualification"))
        self.assertEqual(table.row_count, 3)
        # Numbers alone, even counting up, never make a header.
        self.assertIsNone(detect_header_row([["1", "2"], ["3", "4"]]))
        self.assertIsNone(detect_header_row([["10", "20", "30"], ["11", "21", "31"]]))
        self.assertIsNone(detect_header_row([[1.0, 2.0, 3.0], [4.0, 5.0, 6.0]]))


class TwoRowHeaderTests(unittest.TestCase):
    """Grouped headers on two rows (a marks register) become one header per column."""

    # A note beside the table, longer than any header label.
    LEGEND = "Legend: marks shown in red are below the pass mark set by the university"

    @staticmethod
    def _marks_sheet(subjects: list[str]) -> tuple[list[list[object]], list[str], tuple[str, ...]]:
        # The layout of a result analysis sheet: a title, group labels over the
        # subjects, the subjects, the faculty initials under them, then students.
        count = len(subjects)
        cells = [f"{subject}\n(MB10{number})" for number, subject in enumerate(subjects, start=1)]
        initials = ["PQ", "RS", "TU", "VW", "XY", "ZA"][:count]
        blank = [None] * (count - 1)
        rows: list[list[object]] = [
            [None, "ABC College MBA Result Analysis"],
            ["Sl. No", "USN No", "Name", "EXTERNAL MARKS", *blank, "INTERNAL MARKS", *blank, "Total", "Percentage", None, TwoRowHeaderTests.LEGEND],
            [None, None, None, *cells, *cells, None, None, None, None, "Pass - external filled"],
            [None, None, None, *initials, *initials, None, None, None, None, "Initials of the faculty for each subject"],
        ]
        for number, name in enumerate(["Ravi Kumar", "Asha Rao", "Kiran S"], start=1):
            external = [40 + number + offset for offset in range(count)]
            internal = [15.5 + number + offset for offset in range(count)]
            rows.append([number, f"1MS23MBA00{number}", name, *external, *internal, sum(external) + sum(internal), 60 + number])
        last = 3 + 2 * count  # the column after the subjects (0-based)
        letter = lambda column: chr(65 + column)
        merges = [f"D2:{letter(2 + count)}2", f"{letter(3 + count)}2:{letter(2 + 2 * count)}2", "A2:A4", "B2:B4", "C2:C4", f"{letter(last)}2:{letter(last)}4", f"{letter(last + 1)}2:{letter(last + 1)}4", f"B1:{letter(last + 1)}1"]
        named = [f"{subject} (MB10{number})" for number, subject in enumerate(subjects, start=1)]
        headers = ("Sl. No", "USN No", "Name", *[f"EXTERNAL MARKS {name}" for name in named], *[f"INTERNAL MARKS {name}" for name in named], "Total", "Percentage", f"column_{last + 3}", TwoRowHeaderTests.LEGEND, "Pass - external filled")
        return rows, merges, headers

    def _assert_marks_table(self, table, headers: tuple[str, ...]) -> None:
        self.assertEqual(table.headers, headers)
        self.assertEqual([record.fields["USN No"] for record in table.records], ["1MS23MBA001", "1MS23MBA002", "1MS23MBA003"])
        self.assertEqual([record.row_number for record in table.records], [5, 6, 7])
        self.assertEqual(table.records[0].fields[headers[3]], 41)
        self.assertEqual(table.records[0].fields[headers[3 + (len(headers) - 8) // 2]], 16.5)
        self.assertIn("skipped_1_leading_rows", table.warnings)
        self.assertIn("skipped_row_4_as_annotation:blank USN No, Name; only short labels over number columns", table.warnings)

    def test_grouped_subject_headers_and_the_initials_row(self):
        # Six subjects a group: the subject row is detected and the group row above joins it.
        # Two subjects a group: the group row is detected and the subject row below joins it.
        for subjects in (["Accounts", "Economics", "Law", "Marketing", "Statistics", "Ethics"], ["Accounts", "Economics"]):
            with self.subTest(subjects=len(subjects)):
                rows, merges, headers = self._marks_sheet(subjects)
                self._assert_marks_table(grid_to_table(rows, name="Marks", source_file="marks.xlsx", sheet="Marks"), headers)
                workbook = build_xlsx({"Marks": rows}, merges={"Marks": merges})
                self._assert_marks_table(excel_parser._parse_with_stdlib("marks.xlsx", workbook).tables[0], headers)
                self._assert_marks_table(excel_parser._parse_with_stdlib("marks.xlsx", build_xlsx({"Marks": rows})).tables[0], headers)
                if OPENPYXL_AVAILABLE:
                    self._assert_marks_table(excel_parser._parse_with_openpyxl("marks.xlsx", workbook).tables[0], headers)

    def test_a_lower_header_row_wider_than_the_upper_one(self):
        students = [["1MS23MBA001", "Ravi", 50, 60, 70], ["1MS23MBA002", "Asha", 55, 65, 75]]
        sheets = {
            "under the detected row": [["USN", "Name", "Marks"], [None, None, "Maths", "Physics", "Chem"], *students],
            "over the detected row": [["Class test"], ["USN", "Name", "Marks"], [None, None, "Maths", "Physics", "Chem"], *students],
        }
        for label, rows in sheets.items():
            workbook = build_xlsx({"Marks": rows})
            tables = {"grid": grid_to_table(rows, name="Marks", source_file="marks.xlsx", sheet="Marks"), "stdlib": excel_parser._parse_with_stdlib("marks.xlsx", workbook).tables[0]}
            if OPENPYXL_AVAILABLE:
                tables["openpyxl"] = excel_parser._parse_with_openpyxl("marks.xlsx", workbook).tables[0]
            for reader, table in tables.items():
                with self.subTest(label, reader=reader):
                    self.assertEqual(table.headers, ("USN", "Name", "Marks Maths", "Marks Physics", "Marks Chem"))
                    self.assertEqual([record.fields["USN"] for record in table.records], ["1MS23MBA001", "1MS23MBA002"])
                    self.assertEqual(table.records[1].fields["Marks Chem"], 75)

    def test_merged_cells_decide_how_far_a_group_label_reaches(self):
        rows = [["USN", "Name", "Marks"], [None, None, "Maths", "Physics", "Project"], ["1MS23MBA001", "Ravi", 50, 60, 18], ["1MS23MBA002", "Asha", 55, 65, 19]]
        merged = excel_parser._parse_with_stdlib("marks.xlsx", build_xlsx({"Marks": rows}, merges={"Marks": ["C1:D1"]})).tables[0]
        self.assertEqual(merged.headers, ("USN", "Name", "Marks Maths", "Marks Physics", "Project"))
        # Without merges the label reaches every labelled column up to the next label.
        self.assertEqual(grid_to_table(rows, name="Marks", source_file="marks.csv").headers, ("USN", "Name", "Marks Maths", "Marks Physics", "Marks Project"))
        self.assertEqual(merged.row_count, 2)

    def test_single_row_headers_and_first_data_rows_are_unchanged(self):
        students = [["Ravi Kumar", "MBA", "A"], ["Asha Rao", "MCA", "B"]]
        phones = [["1MS23MBA001", "Ravi", "MBA", "9876543210"], ["1MS23MBA002", "Asha", "MBA", "9876543211"], ["1MS23MBA003", "Kiran", "MCA", "9876543212"]]
        grids = {
            "title and a newline in a header": ([["ABC College"], ["USN", "Name", "Accounts\n(MB101)", "Economics"], ["1MS23MBA001", "Ravi", 45, 50], ["1MS23MBA002", "Asha", 40, 42]], ("USN", "Name", "Accounts (MB101)", "Economics"), 2, ["skipped_1_leading_rows"]),
            "key/value lines above": ([["Class", "MBA III Sem"], ["Faculty", "Dr. Rao"], ["Name", "Program", "Section"], *students], ("Name", "Program", "Section"), 2, ["skipped_2_leading_rows"]),
            "a spread-out title over a complete header": ([["ABC College Student List", "", "", "Date: 01-01-2025"], ["USN", "Name", "Program", "Section", "Phone"], ["1MS23MBA001", "Ravi", "MBA", "A", "9876543210"], ["1MS23MBA002", "Asha", "MBA", "B", "9876543211"]], ("USN", "Name", "Program", "Section", "Phone"), 2, ["skipped_1_leading_rows"]),
            "a first row with no student": ([["USN", "Name", "M1", "M2", "M3", "M4", "M5", "Transport", ""], ["", "", "", "", "", "", "", "Bus", "Route 5"], ["1MS23MBA002", "Asha", 40, 41, 42, 43, 44, "Van", "Route 2"], ["1MS23MBA003", "Kiran", 30, 31, 32, 33, 34, "Bus", "Route 7"]], ("USN", "Name", "M1", "M2", "M3", "M4", "M5", "Transport", "column_9"), 3, []),
            "a first student absent in every subject": ([["USN", "Name", "Maths", "Physics"], ["1MS23MBA001", "Ravi", "AB", "AB"], ["1MS23MBA002", "Asha", 40, 42], ["1MS23MBA003", "Kiran", 41, 43]], ("USN", "Name", "Maths", "Physics"), 3, []),
            "a group line over a one-column label": ([["Faculty Details", "", "Department Info"], ["Name", "Designation", "Department"], ["Ravi Kumar", "Professor", "Marketing"], ["Asha Rao", "Associate Professor", "Finance"]], ("Name", "Designation", "Department"), 2, ["skipped_1_leading_rows"]),
            # A title or key/value line never joins a header that leaves a cell blank.
            "a title and a date over a header with a blank cell": ([["ABC College Student List", "", "", "Date: 01-01-2025"], ["USN", "Name", "Program", ""], *phones], ("USN", "Name", "Program", "column_4"), 3, ["skipped_1_leading_rows"]),
            "a key/value line over a header with a blank cell": ([["Class", "", "", "MBA III"], ["USN", "Name", "Program", ""], *phones], ("USN", "Name", "Program", "column_4"), 3, ["skipped_1_leading_rows"]),
            "a department line over a header with a blank cell": ([["Department: MBA", "", "", "AY 2024-25"], ["USN", "Name", "Program", ""], *phones], ("USN", "Name", "Program", "column_4"), 3, ["skipped_1_leading_rows"]),
            "a staff title over a header with a blank cell": ([["Staff Details", "", "", "Updated on 01-Jan-2025"], ["Name", "Designation", "Department", ""], ["Ravi Kumar", "Professor", "Marketing", "9876543210"], ["Asha Rao", "Associate Professor", "Finance", "9876543211"]], ("Name", "Designation", "Department", "column_4"), 2, ["skipped_1_leading_rows"]),
            # A first student with no id and absence codes for marks is a student, not faculty initials.
            "a first student absent with no USN": ([["USN", "IA1", "IA2"], ["", "AB", "AB"], ["1MS23MBA002", 20, 21], ["1MS23MBA003", 22, 23]], ("USN", "IA1", "IA2"), 3, []),
            "a first student absent with no USN, codes with dots": ([["USN", "IA1", "IA2"], ["", "Abs.", "N.A."], ["1MS23MBA002", 20, 21], ["1MS23MBA003", 22, 23]], ("USN", "IA1", "IA2"), 3, []),
            "a first student absent with no roll number or name": ([["Roll", "Name", "IA1", "IA2", "IA3"], ["", "", "AB", "NE", "-"], [2, "Asha", 20, 21, 19], [3, "Kiran", 22, 23, 18]], ("Roll", "Name", "IA1", "IA2", "IA3"), 3, []),
            "a first student absent with a name": ([["USN", "Name", "IA1", "IA2"], ["", "Ravi Kumar", "AB", "A"], ["1MS23MBA002", "Asha", 20, 21], ["1MS23MBA003", "Kiran", 22, 23]], ("USN", "Name", "IA1", "IA2"), 3, []),
        }
        for label, (grid, headers, count, warnings) in grids.items():
            with self.subTest(label):
                table = grid_to_table(grid, name="t", source_file="t.xlsx", sheet="t")
                self.assertEqual(table.headers, headers)
                self.assertEqual(table.row_count, count)
                self.assertEqual(table.warnings, warnings)

    def test_a_merged_title_never_joins_a_header_with_a_blank_cell(self):
        rows = [["ABC College Student List", None, None, "Date: 01-01-2025"], ["USN", "Name", "Program"], ["1MS23MBA001", "Ravi", "MBA", "9876543210"], ["1MS23MBA002", "Asha", "MBA", "9876543211"]]
        workbook = build_xlsx({"Students": rows}, merges={"Students": ["A1:C1"]})
        tables = {"stdlib": excel_parser._parse_with_stdlib("students.xlsx", workbook).tables[0]}
        if OPENPYXL_AVAILABLE:
            tables["openpyxl"] = excel_parser._parse_with_openpyxl("students.xlsx", workbook).tables[0]
        for reader, table in tables.items():
            with self.subTest(reader=reader):
                self.assertEqual(table.headers, ("USN", "Name", "Program", "column_4"))
                self.assertEqual([record.fields["USN"] for record in table.records], ["1MS23MBA001", "1MS23MBA002"])
                self.assertEqual(table.warnings, ["skipped_1_leading_rows"])

    def test_the_same_subjects_under_two_group_labels_join_one_identifier(self):
        rows = [["Class test"], ["USN", "IA1", None, "IA2", None], [None, "Maths", "Physics", "Maths", "Physics"], ["1MS23MBA001", 50, 60, 70, 80], ["1MS23MBA002", 55, 65, 75, 85]]
        table = grid_to_table(rows, name="Marks", source_file="marks.xlsx", sheet="Marks")
        self.assertEqual(table.headers, ("USN", "IA1 Maths", "IA1 Physics", "IA2 Maths", "IA2 Physics"))
        self.assertEqual(table.warnings, ["skipped_1_leading_rows", "headers_combined_from_rows_2_and_3"])
        self.assertEqual(table.records[1].fields["IA2 Physics"], 85)

    @staticmethod
    def _read_with_every_reader(rows: list[list[object]], merges: list[str]) -> dict[str, object]:
        workbook = build_xlsx({"Sheet": rows}, merges={"Sheet": merges})
        tables = {"grid": grid_to_table(rows, name="Sheet", source_file="sheet.xlsx", sheet="Sheet"), "stdlib": excel_parser._parse_with_stdlib("sheet.xlsx", workbook).tables[0]}
        if OPENPYXL_AVAILABLE:
            tables["openpyxl"] = excel_parser._parse_with_openpyxl("sheet.xlsx", workbook).tables[0]
        return tables

    def test_one_identifier_column_joins_a_single_group_label(self):
        # "USN" merged over both header rows, then one "Marks" group over the subjects.
        subjects = ["Maths", "Physics", "Chemistry", "Biology", "English", "Kannada"]
        students = ["1MS23MBA001", "1MS23MBA002", "1MS23MBA003"]
        layouts = {
            "USN": ([["USN", "Marks", *[None] * 5], [None, *subjects], *[[usn, 40 + row, 41, 42, 43, 44, 45] for row, usn in enumerate(students)]], ["A1:A2", "B1:G1"], ("USN",)),
            "Sl. No and Hall Ticket No": ([["Sl. No", "Hall Ticket No", "Marks", *[None] * 5], [None, None, *subjects], *[[row + 1, usn, 40 + row, 41, 42, 43, 44, 45] for row, usn in enumerate(students)]], ["A1:A2", "B1:B2", "C1:H1"], ("Sl. No", "Hall Ticket No")),
        }
        for label, (rows, merges, identifiers) in layouts.items():
            for reader, table in self._read_with_every_reader(rows, merges).items():
                with self.subTest(label, reader=reader):
                    self.assertEqual(table.headers, (*identifiers, *[f"Marks {subject}" for subject in subjects]))
                    self.assertEqual(table.warnings, ["headers_combined_from_rows_1_and_2"])
                    self.assertEqual([record.fields[identifiers[-1]] for record in table.records], students)
                    self.assertEqual(table.records[2].fields["Marks Maths"], 42)

    def test_a_key_value_line_with_a_name_never_joins_a_header_with_blank_cells(self):
        # The key or value holds "Name" and two data columns have no header of their own.
        header = ["USN", "Name", "Program", "Section", "Batch", "Gender", "Category", "Quota"]
        students = [[f"1MS23MBA00{row}", name, "MBA", "A", "2024", "M", "GM", "CET", f"987654321{row}", f"s{row}@example.edu"] for row, name in enumerate(["Ravi Kumar", "Asha Rao", "Kiran S"], start=1)]
        for above in (["Class: MBA III", *[None] * 7, "Faculty Name", "Dr. Rao"], ["Department of Commerce", *[None] * 7, "Name of HOD", "Dr. Rao"]):
            for reader, table in self._read_with_every_reader([above, header, *students], []).items():
                with self.subTest(above[0], reader=reader):
                    self.assertEqual(table.headers, (*header, "column_9", "column_10"))
                    self.assertEqual(table.warnings, ["skipped_1_leading_rows"])
                    self.assertEqual([record.fields["USN"] for record in table.records], ["1MS23MBA001", "1MS23MBA002", "1MS23MBA003"])

    def test_faculty_initials_that_read_as_absence_codes_under_a_grouped_header(self):
        subjects = ["Accounts (MB101)", "Economics (MB102)"]
        rows = [
            [None, "Result"],
            ["Sl. No", "USN No", "Name", "EXTERNAL MARKS", None, "INTERNAL MARKS", None],
            [None, None, None, *subjects, *subjects],
            [None, None, None, "AB", "MP", "AB", "MP"],
            *[[row, f"1MS23MBA00{row}", name, 40 + row, 42, 16, 17] for row, name in enumerate(["Ravi Kumar", "Asha Rao", "Kiran S"], start=1)],
        ]
        for reader, table in self._read_with_every_reader(rows, ["A2:A4", "B2:B4", "C2:C4", "D2:E2", "F2:G2"]).items():
            with self.subTest(reader=reader):
                self.assertEqual(table.headers, ("Sl. No", "USN No", "Name", *[f"EXTERNAL MARKS {subject}" for subject in subjects], *[f"INTERNAL MARKS {subject}" for subject in subjects]))
                self.assertEqual([record.row_number for record in table.records], [5, 6, 7])
                self.assertIn("skipped_row_4_as_annotation:blank USN No, Name; only short labels over number columns", table.warnings)


class WideSheetTests(unittest.TestCase):
    def test_giant_cell_references_and_wide_rows_are_bounded(self):
        import time

        from app.ingestion.parsers.excel_parser import MAX_SHEET_COLUMNS, parse_xlsx

        content = build_xlsx({"Sheet1": [["Name", "USN"], ["Ravi", "MBA001"]]})
        wide = "".join(f'<c r="{chr(65 + (i % 26))}{chr(65 + (i // 26 % 26))}{chr(65 + (i // 676 % 26))}3"><v>{i}</v></c>' for i in range(3000))

        def transform(data: bytes) -> bytes:
            text = data.decode("utf-8")
            text = text.replace("</row></sheetData>", '<c r="ZZZZZZ2" t="inlineStr"><is><t>far</t></is></c></row><row r="3">' + wide + "</row></sheetData>")
            return text.encode("utf-8")

        crafted = rewrite_entry(content, "xl/worksheets/sheet1.xml", transform)
        started = time.perf_counter()
        result = parse_xlsx("wide.xlsx", crafted)
        self.assertLess(time.perf_counter() - started, 5.0)
        table = result.tables[0]
        self.assertLessEqual(len(table.headers), MAX_SHEET_COLUMNS)
        self.assertLessEqual(max(len(record.fields) for record in table.records), MAX_SHEET_COLUMNS)
        self.assertNotIn("far", str([record.fields for record in table.records]), "a cell beyond the last supported column is ignored, never padded out to")


if __name__ == "__main__":
    unittest.main()
